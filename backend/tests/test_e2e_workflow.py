"""Phase 16: scripted end-to-end workflow test.

Exercises the full real path through FastAPI's TestClient against the real (temp sqlite) test
DB: admin login -> load demo gearbox project -> fetch analysis/components -> generate BOM ->
edit a line item -> fetch BOM tree (rollup correct) -> procurement list (derived, not
fabricated) -> approve BOM (role-gated: viewer rejected, admin allowed) -> export CSV/Excel
(flat + hierarchical) with real persisted values checked -> audit log entry for the approval
with the real admin identity.

Also covers a real-upload-path test: a small synthetic PDF (generated on the fly, never
committed as a binary) through the full real pipeline with the AI provider mocked (no API key
in CI), asserting it is never demo-tagged and produces its own independent analysis record.
"""
from __future__ import annotations

import io
import os

import openpyxl
import pandas as pd
import pytest
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.config import get_settings
from app.models.models import AuditLog, Bom, Drawing, DrawingStatus, MakeBuy, UserRole
from app.schemas.analysis import AnalysisResult, ComponentCandidate, TitleBlock


def _login(client, email, password):
    resp = client.post("/api/auth/login", json={"email": email, "password": password})
    assert resp.status_code == 200, resp.text
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def test_full_gearbox_workflow_end_to_end(client, make_user, db):
    from fastapi.testclient import TestClient
    from app.main import app

    admin, admin_pw = make_user(role=UserRole.ADMIN)
    viewer, viewer_pw = make_user(role=UserRole.VIEWER)
    # Separate TestClient instances (and so separate cookie jars) for each user - sharing one
    # client would let the second login's httpOnly cookie silently override the first's and
    # mask the role-gating being tested here.
    admin_client = TestClient(app)
    viewer_client = TestClient(app)
    admin_headers = _login(admin_client, admin.email, admin_pw)
    viewer_headers = _login(viewer_client, viewer.email, viewer_pw)
    client = admin_client

    # 2. Load the demo gearbox project via the existing demo-load endpoint.
    load_resp = client.post("/api/demo/load", headers=admin_headers)
    assert load_resp.status_code == 200, load_resp.text

    gearbox = db.query(Drawing).filter(Drawing.drawing_number == "GB-1001").first()
    assert gearbox is not None
    assert gearbox.is_demo is True

    # 3. Fetch its analysis/components, confirm evidence/confidence fields are present.
    analyses_resp = client.get(f"/api/drawings/{gearbox.id}/analyses", headers=admin_headers)
    assert analyses_resp.status_code == 200
    analyses = analyses_resp.json()
    assert len(analyses) >= 1
    analysis = analyses[0]
    analysis_detail = client.get(f"/api/analyses/{analysis['id']}", headers=admin_headers).json()
    assert analysis_detail["components"], "expected seeded gearbox components"
    for comp in analysis_detail["components"]:
        assert "confidence" in comp and comp["confidence"] is not None
        assert "source_text" in comp or "source_page" in comp  # evidence fields present

    # 4. Generate a BOM from it.
    gen_resp = client.post(f"/api/analyses/{analysis['id']}/generate-bom", headers=admin_headers)
    assert gen_resp.status_code == 200, gen_resp.text
    bom = gen_resp.json()
    assert bom["line_items"], "BOM generation should produce line items"

    # 5. Edit a BOM line item (change quantity).
    target_line = bom["line_items"][0]
    new_qty = (target_line["quantity"] or 1) + 7
    patch_resp = client.patch(
        f"/api/boms/line-items/{target_line['id']}", json={"quantity": new_qty}, headers=admin_headers
    )
    assert patch_resp.status_code == 200, patch_resp.text

    # 6. Fetch the BOM tree endpoint, confirm the edit is reflected and roll-up quantities correct.
    tree_resp = client.get(f"/api/boms/{bom['id']}/tree", headers=admin_headers)
    assert tree_resp.status_code == 200
    tree = tree_resp.json()

    def find_node(nodes, line_id):
        for n in nodes:
            if n.get("id") == line_id:
                return n
            found = find_node(n.get("children", []), line_id)
            if found:
                return found
        return None

    node = find_node(tree["nodes"], target_line["id"])
    assert node is not None
    assert node["quantity"] == new_qty

    db_bom = db.query(Bom).get(bom["id"])
    db_line = next(li for li in db_bom.line_items if li.id == target_line["id"])
    assert db_line.quantity == new_qty
    # Roll-up: total_quantity for a top-level item with no parent equals its own quantity.
    if db_line.parent_line_item_id is None:
        assert db_line.total_quantity == new_qty

    # 7. Generate/fetch the procurement list, confirm Buy/Make/Review classification is present
    # and derived from the BOM (not fabricated) - cross-check against the persisted BOM items.
    proc_resp = client.get(f"/api/procurement/by-drawing/{gearbox.id}", headers=admin_headers)
    assert proc_resp.status_code == 200
    groups = proc_resp.json()
    assert set(groups.keys()) >= {"buy", "make", "review"}
    all_proc_ids = {item["id"] for bucket in groups.values() for item in bucket}
    all_bom_ids = {li.id for li in db_bom.line_items}
    assert all_proc_ids == all_bom_ids
    for li in db_bom.line_items:
        key = (li.make_buy.value if hasattr(li.make_buy, "value") else li.make_buy).lower()
        bucket_ids = {item["id"] for item in groups[key]}
        assert li.id in bucket_ids

    # 8. Approve the BOM (role-gated): viewer rejected, admin succeeds.
    viewer_approve = viewer_client.post(f"/api/boms/{bom['id']}/approve", headers=viewer_headers)
    assert viewer_approve.status_code == 403
    admin_approve = client.post(f"/api/boms/{bom['id']}/approve", headers=admin_headers)
    assert admin_approve.status_code == 200, admin_approve.text
    db.refresh(db_bom)
    assert db_bom.status == "Approved"

    # 9. Export CSV and Excel (flat + hierarchical), assert exported values match persisted data.
    csv_resp = client.get(f"/api/boms/{bom['id']}/export/bom.csv", headers=admin_headers)
    assert csv_resp.status_code == 200
    csv_df = pd.read_csv(io.BytesIO(csv_resp.content))
    exported_qty = csv_df.loc[csv_df["Part Number"] == target_line["part_number"], "Quantity"].iloc[0]
    assert float(exported_qty) == float(new_qty)
    assert "Level" in csv_df.columns  # hierarchical (level/indent) info present in the flat export

    xlsx_resp = client.get(f"/api/boms/{bom['id']}/export/bom.xlsx", headers=admin_headers)
    assert xlsx_resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
    ws = wb["Sheet1"]
    headers_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    qty_col = headers_row.index("Quantity")
    pn_col = headers_row.index("Part Number")
    found = False
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[pn_col] == target_line["part_number"]:
            assert float(row[qty_col]) == float(new_qty)
            found = True
    assert found, "expected the edited line item's real quantity in the exported Excel file"

    purchase_csv = client.get(f"/api/boms/{bom['id']}/export/purchase-list.csv", headers=admin_headers)
    assert purchase_csv.status_code == 200
    purchase_df = pd.read_csv(io.BytesIO(purchase_csv.content))
    assert set(purchase_df["Make/Buy"].unique()) <= {"Buy", "Review"}

    # 10. Assert an audit log entry exists for the approval action with the real admin identity.
    audit = (
        db.query(AuditLog)
        .filter(AuditLog.entity_type == "bom", AuditLog.entity_id == bom["id"], AuditLog.action == "approved")
        .order_by(AuditLog.created_at.desc())
        .first()
    )
    assert audit is not None
    assert audit.actor == admin.email


def _synthetic_pdf_bytes(text: str) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.setFont("Helvetica", 18)
    c.drawString(72, 700, text)
    c.showPage()
    c.save()
    return buf.getvalue()


def test_real_pdf_upload_full_path_mocked_provider_never_demo_tagged(client, make_user, db, monkeypatch):
    """Uploads a small synthetic PDF through the real /api/drawings upload + analyze endpoints,
    with the AI provider mocked (no API key available in CI). Confirms the resulting Drawing and
    Analysis are never tagged/contaminated with demo data and are independent of the seeded
    demo projects."""
    admin, admin_pw = make_user(role=UserRole.ADMIN)
    headers = _login(client, admin.email, admin_pw)

    pdf_bytes = _synthetic_pdf_bytes("E2E-REAL-UPLOAD-MARKER")

    class MockedProvider:
        name = "openai"

        def analyze(self, *, image_path, ocr_text, drawing_number):
            return AnalysisResult(
                provider="openai",
                title_block=TitleBlock(drawing_number=drawing_number, title="Real Uploaded Part",
                                        revision="A", units="mm", scale="1:1", material_spec="Steel"),
                dimensions=[],
                materials=["Steel"],
                components=[
                    ComponentCandidate(
                        part_number="REAL-UP-1", description="Real uploaded bracket", category="Bracket",
                        quantity=1, material="Steel", make_buy="Make", confidence=0.9,
                        verification_status="Unverified", source_page=1, source_text="mocked vision result",
                    )
                ],
                raw_ocr_text=ocr_text,
                overall_confidence=0.9,
            )

    monkeypatch.setattr("app.services.analysis_pipeline.get_provider", lambda: MockedProvider())

    upload_resp = client.post(
        "/api/drawings",
        headers=headers,
        files={"file": ("e2e_real_upload.pdf", pdf_bytes, "application/pdf")},
        data={"drawing_number": "E2E-REAL-1", "title": "E2E Real Upload", "revision": "A"},
    )
    assert upload_resp.status_code == 200, upload_resp.text
    drawing_id = upload_resp.json()["id"]
    assert upload_resp.json()["is_demo"] is False

    analyze_resp = client.post(f"/api/drawings/{drawing_id}/analyze", headers=headers)
    assert analyze_resp.status_code == 200, analyze_resp.text
    analysis = analyze_resp.json()

    drawing = db.query(Drawing).get(drawing_id)
    assert drawing.is_demo is False
    assert drawing.status in (DrawingStatus.NEEDS_REVIEW, DrawingStatus.REVIEWED)
    assert analysis["provider"] == "openai"
    assert any(c["part_number"] == "REAL-UP-1" for c in analysis["components"])

    # Independent of seeded demo content: none of the demo project drawing numbers/part numbers
    # appear in this analysis's output.
    from seed import PROJECTS
    seeded_numbers = {p["drawing_number"] for p in PROJECTS}
    seeded_parts = {c["part_number"] for p in PROJECTS for c in p["components"]}
    assert analysis["title_block"]["drawing_number"] not in seeded_numbers
    assert all(c["part_number"] not in seeded_parts for c in analysis["components"])
